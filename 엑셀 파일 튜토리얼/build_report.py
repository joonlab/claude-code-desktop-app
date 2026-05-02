"""dummy_sales.csv를 읽어 미니멀한 스타일의 엑셀 보고서를 생성한다.

색상 팔레트
- Red             #A50034  (헤드라인 강조 / 누적바)
- Neutral Gray    #6B6B6B  (본문 텍스트)
- Premium Silver  #8A8D8F  (헤더 라벨, 보조선)
- Highlight Gold  #85714D  (KPI 강조값)
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawFont,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RED = "A50034"
NEUTRAL = "6B6B6B"
SILVER = "8A8D8F"
GOLD = "85714D"
INK = "1F1F1F"
BAND = "FAFAFA"

FONT = "Calibri"

HERE = Path(__file__).parent
SRC = HERE / "dummy_sales.csv"
OUT = HERE / "sales_report.xlsx"


def thin_bottom(color: str) -> Border:
    return Border(bottom=Side(style="thin", color=color))


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def label_cell(ws: Worksheet, coord: str, text: str) -> None:
    cell = ws[coord]
    cell.value = text.upper()
    cell.font = Font(name=FONT, size=9, bold=True, color=SILVER)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def kpi_cell(ws: Worksheet, coord: str, text: str, color: str) -> None:
    cell = ws[coord]
    cell.value = text
    cell.font = Font(name=FONT, size=18, bold=True, color=color)
    cell.alignment = Alignment(horizontal="left", vertical="center")


# ---------------------------------------------------------------------------
# Sheet 1: 보고서 개요
# ---------------------------------------------------------------------------


def build_overview(ws: Worksheet, df: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 3, "B": 22, "C": 22, "D": 22, "E": 22, "F": 4})

    ws["B2"] = "2025 ANNUAL SALES REVIEW"
    ws["B2"].font = Font(name=FONT, size=10, bold=True, color=GOLD)

    ws.merge_cells("B3:E3")
    ws["B3"] = "프리미엄 가전 매출 보고서"
    ws["B3"].font = Font(name=FONT, size=26, bold=True, color=INK)
    ws.row_dimensions[3].height = 36

    ws.merge_cells("B4:E4")
    ws["B4"] = "Premium Electronics · Annual Performance Report"
    ws["B4"].font = Font(name=FONT, size=11, color=NEUTRAL, italic=True)

    ws.row_dimensions[6].height = 4
    for col in ("B", "C", "D", "E"):
        ws[f"{col}6"].fill = PatternFill("solid", fgColor=RED)

    total_rev = int(df["매출"].sum())
    total_units = int(df["판매수량"].sum())
    avg_order = int(df["매출"].mean())
    top_region = df.groupby("지역")["매출"].sum().idxmax()

    label_cell(ws, "B9", "총 매출")
    kpi_cell(ws, "B10", f"₩{total_rev:,}", GOLD)

    label_cell(ws, "C9", "총 판매 수량")
    kpi_cell(ws, "C10", f"{total_units:,} 대", INK)

    label_cell(ws, "D9", "평균 주문 매출")
    kpi_cell(ws, "D10", f"₩{avg_order:,}", NEUTRAL)

    label_cell(ws, "E9", "최고 매출 지역")
    kpi_cell(ws, "E10", str(top_region), RED)

    ws.row_dimensions[9].height = 16
    ws.row_dimensions[10].height = 30

    for col in ("B", "C", "D", "E"):
        ws[f"{col}11"].border = thin_bottom(SILVER)

    ws.merge_cells("B14:E14")
    ws["B14"] = "보고서 구성"
    ws["B14"].font = Font(name=FONT, size=11, bold=True, color=INK)

    sections = [
        ("01", "원본 데이터", "180건의 트랜잭션 원본을 정렬된 표로 제공"),
        ("02", "분석 결과", "월별 매출 추이 · 지역/카테고리 분포 · 담당자 랭킹"),
    ]
    base_row = 16
    for i, (idx, title, desc) in enumerate(sections):
        row = base_row + i * 2
        ws[f"B{row}"] = idx
        ws[f"B{row}"].font = Font(name=FONT, size=10, bold=True, color=GOLD)
        ws[f"C{row}"] = title
        ws[f"C{row}"].font = Font(name=FONT, size=11, bold=True, color=INK)
        ws.merge_cells(f"D{row}:E{row}")
        ws[f"D{row}"] = desc
        ws[f"D{row}"].font = Font(name=FONT, size=10, color=NEUTRAL)
        ws.row_dimensions[row].height = 18


# ---------------------------------------------------------------------------
# Sheet 2: 원본 데이터
# ---------------------------------------------------------------------------


def build_raw(ws: Worksheet, df: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    widths = {"A": 3, "B": 13, "C": 10, "D": 12, "E": 22, "F": 12, "G": 14, "H": 16, "I": 12}
    set_widths(ws, widths)

    ws["B2"] = "원본 트랜잭션"
    ws["B2"].font = Font(name=FONT, size=18, bold=True, color=INK)
    ws["B3"] = "RAW SALES DATA · 2025"
    ws["B3"].font = Font(name=FONT, size=9, bold=True, color=GOLD)

    headers = list(df.columns)
    header_row = 5
    for i, h in enumerate(headers):
        cell = ws.cell(row=header_row, column=2 + i, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color=SILVER)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_bottom(SILVER)
    ws.row_dimensions[header_row].height = 22

    money_cols = {"단가", "매출"}
    int_cols = {"판매수량"}

    for r_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        fill = PatternFill("solid", fgColor=BAND) if r_idx % 2 == 0 else None
        for c_idx, (col_name, value) in enumerate(zip(headers, row), start=2):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = Font(name=FONT, size=10, color=NEUTRAL)
            if col_name in money_cols:
                cell.number_format = '"₩"#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in int_cols:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if fill:
                cell.fill = fill

    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)


# ---------------------------------------------------------------------------
# Sheet 3: 분석 결과
# ---------------------------------------------------------------------------


def section_title_at(ws: Worksheet, row: int, col: int, text: str, sub: str) -> None:
    ws.cell(row=row, column=col, value=text).font = Font(
        name=FONT, size=14, bold=True, color=INK
    )
    ws.cell(row=row + 1, column=col, value=sub).font = Font(
        name=FONT, size=9, bold=True, color=GOLD
    )


def section_title(ws: Worksheet, row: int, text: str, sub: str) -> None:
    section_title_at(ws, row, 2, text, sub)


def write_table(
    ws: Worksheet,
    start_row: int,
    start_col: int,
    df: pd.DataFrame,
    money_cols: set[str] | None = None,
    int_cols: set[str] | None = None,
    highlight_top: bool = False,
) -> tuple[int, int]:
    money_cols = money_cols or set()
    int_cols = int_cols or set()

    for i, col in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + i, value=col)
        cell.font = Font(name=FONT, size=10, bold=True, color=SILVER)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_bottom(SILVER)
    ws.row_dimensions[start_row].height = 20

    max_value = df[df.columns[-1]].max() if highlight_top else None

    for r, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c, (col_name, value) in enumerate(zip(df.columns, row), start=start_col):
            cell = ws.cell(row=r, column=c, value=value)
            is_top = highlight_top and col_name == df.columns[-1] and value == max_value
            color = GOLD if is_top else NEUTRAL
            bold = bool(is_top)
            cell.font = Font(name=FONT, size=10, color=color, bold=bold)
            if col_name in money_cols:
                cell.number_format = '"₩"#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in int_cols:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    end_row = start_row + len(df)
    end_col = start_col + len(df.columns) - 1
    return end_row, end_col


def chart_text(color: str = NEUTRAL, size_pt: int = 9, bold: bool = False) -> RichText:
    cp = CharacterProperties(
        sz=size_pt * 100,
        b=bold,
        solidFill=color,
        latin=DrawFont(typeface=FONT),
    )
    return RichText(
        bodyPr=RichTextProperties(),
        p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp, r=[])],
    )


def style_chart(
    chart,
    *,
    title: str | None = None,
    value_format: str = '#,##0,,"M"',
    show_value_labels: bool = False,
) -> None:
    chart.height = 9
    chart.width = 18
    chart.legend = None

    if title:
        chart.title = title

    for axis in (chart.x_axis, chart.y_axis):
        if axis is None:
            continue
        axis.delete = False
        axis.majorGridlines = None
        axis.txPr = chart_text(NEUTRAL, 9)

    # In openpyxl, chart.y_axis is always the ValueAxis (regardless of bar orientation).
    if chart.y_axis is not None:
        chart.y_axis.number_format = value_format

    if show_value_labels:
        chart.dLbls = DataLabelList(
            showVal=True,
            showCatName=False,
            showSerName=False,
            showLegendKey=False,
        )
        chart.dLbls.numFmt = value_format
        chart.dLbls.txPr = chart_text(NEUTRAL, 8)


def build_analysis(ws: Worksheet, df: pd.DataFrame) -> None:
    ws.sheet_view.showGridLines = False
    set_widths(
        ws,
        {"A": 3, "B": 18, "C": 18, "D": 18, "E": 18, "F": 4, "G": 18, "H": 18, "I": 18},
    )

    ws["B2"] = "분석 결과"
    ws["B2"].font = Font(name=FONT, size=20, bold=True, color=INK)
    ws["B3"] = "ANALYSIS · KEY BREAKDOWNS"
    ws["B3"].font = Font(name=FONT, size=9, bold=True, color=GOLD)

    # 1) 월별 매출
    df_month = (
        df.assign(월=pd.to_datetime(df["날짜"]).dt.month)
        .groupby("월", as_index=False)["매출"]
        .sum()
        .rename(columns={"매출": "월매출"})
    )
    df_month["월"] = df_month["월"].apply(lambda m: f"{m:02d}월")

    section_title(ws, 5, "월별 매출 추이", "MONTHLY REVENUE")
    end_row, end_col = write_table(
        ws, 8, 2, df_month, money_cols={"월매출"}
    )

    line = LineChart()
    data_ref = Reference(ws, min_col=end_col, min_row=8, max_row=end_row)
    cat_ref = Reference(ws, min_col=2, min_row=9, max_row=end_row)
    line.add_data(data_ref, titles_from_data=True)
    line.set_categories(cat_ref)
    series = line.series[0]
    series.graphicalProperties.line.solidFill = RED
    series.graphicalProperties.line.width = 18000
    series.smooth = False
    series.marker = Marker(symbol="circle", size=6)
    series.marker.graphicalProperties.solidFill = RED
    series.marker.graphicalProperties.line.solidFill = RED
    style_chart(line, title="월별 매출 (단위: 백만원)")
    ws.add_chart(line, "G5")

    # 2) 지역별 매출
    df_region = (
        df.groupby("지역", as_index=False)["매출"]
        .sum()
        .sort_values("매출", ascending=False)
        .rename(columns={"매출": "지역매출"})
    )

    base = end_row + 4
    section_title(ws, base, "지역별 매출", "REVENUE BY REGION")
    region_end_row, region_end_col = write_table(
        ws, base + 3, 2, df_region, money_cols={"지역매출"}, highlight_top=True
    )

    bar = BarChart()
    bar.type = "bar"
    bar.style = 2
    data_ref = Reference(ws, min_col=region_end_col, min_row=base + 3, max_row=region_end_row)
    cat_ref = Reference(ws, min_col=2, min_row=base + 4, max_row=region_end_row)
    bar.add_data(data_ref, titles_from_data=True)
    bar.set_categories(cat_ref)
    bar.series[0].graphicalProperties.solidFill = RED
    bar.series[0].graphicalProperties.line.solidFill = RED
    style_chart(bar, title="지역별 매출 (단위: 백만원)", show_value_labels=True)
    ws.add_chart(bar, f"G{base}")

    # 3) 카테고리별 매출
    df_cat = (
        df.groupby("카테고리", as_index=False)
        .agg(판매수량=("판매수량", "sum"), 매출=("매출", "sum"))
        .sort_values("매출", ascending=False)
    )
    base2 = region_end_row + 5
    section_title_at(ws, base2, 2, "카테고리별 실적", "REVENUE BY CATEGORY")
    write_table(
        ws,
        base2 + 3,
        2,
        df_cat,
        money_cols={"매출"},
        int_cols={"판매수량"},
        highlight_top=True,
    )

    # 4) Top 담당자 (카테고리 섹션과 같은 행, 우측 칼럼)
    df_rep = (
        df.groupby("담당자", as_index=False)["매출"]
        .sum()
        .sort_values("매출", ascending=False)
        .head(5)
        .rename(columns={"매출": "총매출"})
    )
    section_title_at(ws, base2, 7, "Top 5 담당자", "TOP SALES REPS")
    write_table(ws, base2 + 3, 7, df_rep, money_cols={"총매출"}, highlight_top=True)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def main() -> Path:
    df = pd.read_csv(SRC)

    wb = Workbook()
    overview = wb.active
    overview.title = "개요"
    raw = wb.create_sheet("원본 데이터")
    analysis = wb.create_sheet("분석 결과")

    build_overview(overview, df)
    build_raw(raw, df)
    build_analysis(analysis, df)

    wb.save(OUT)
    print(f"보고서 저장 완료 -> {OUT}")
    return OUT


if __name__ == "__main__":
    main()
